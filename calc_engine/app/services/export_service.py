"""计算书 Word 导出服务"""

from io import BytesIO
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
import json
from datetime import datetime


def build_calculation_report(
    module_name: str,
    module_version: str,
    status: str,
    input_json: str,
    output_json: str,
    references: list = None,
) -> BytesIO:
    """生成计算书 Word 文档"""
    doc = Document()

    # 设置页面边距
    for section in doc.sections:
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)

    # ========== 封面标题 ==========
    title = doc.add_heading('实验室工程计算书', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(f'模块：{module_name} v{module_version}')
    run.font.size = Pt(14)

    date_para = doc.add_paragraph()
    date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_run = date_para.add_run(f'生成时间：{datetime.now().strftime("%Y-%m-%d %H:%M:%S")}')
    date_run.font.size = Pt(10)

    doc.add_paragraph('')
    doc.add_paragraph('')

    # ========== 一、计算状态 ==========
    doc.add_heading('一、计算状态', level=1)
    status_map = {"success": "✅ 计算成功", "warning": "⚠️ 有警告", "error": "❌ 计算失败"}
    doc.add_paragraph(status_map.get(status, status))

    # ========== 二、输入参数 ==========
    doc.add_heading('二、输入参数', level=1)
    try:
        input_data = json.loads(input_json) if isinstance(input_json, str) else input_json
        _write_json_to_doc(doc, input_data)
    except Exception:
        doc.add_paragraph(str(input_json))

    # ========== 三、计算结果 ==========
    doc.add_heading('三、计算结果', level=1)
    try:
        output_data = json.loads(output_json) if isinstance(output_json, str) else output_json
        _write_json_to_doc(doc, output_data)
    except Exception:
        doc.add_paragraph(str(output_json))

    # ========== 四、规范依据 ==========
    if references:
        doc.add_heading('四、规范依据', level=1)
        for ref in references:
            doc.add_paragraph(ref, style='List Bullet')

    # ========== 五、免责声明 ==========
    doc.add_heading('五、免责声明', level=1)
    doc.add_paragraph('本计算书由系统自动生成，计算结果需经持证工程师复核确认后方可用于施工。')
    doc.add_paragraph('引用规范以最新有效版本为准。')

    # 保存到 BytesIO
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def _write_json_to_doc(doc: Document, data, level: int = 0):
    """递归把 JSON 数据写入 Word 文档"""
    if isinstance(data, dict):
        for key, value in data.items():
            if isinstance(value, (dict, list)):
                doc.add_paragraph(f'{key}：')
                _write_json_to_doc(doc, value, level + 1)
            else:
                doc.add_paragraph(f'{key}：{value}', style='List Bullet' if level > 0 else None)
    elif isinstance(data, list):
        for item in data:
            if isinstance(item, (dict, list)):
                _write_json_to_doc(doc, item, level + 1)
            else:
                doc.add_paragraph(str(item), style='List Bullet')
    else:
        doc.add_paragraph(str(data))
